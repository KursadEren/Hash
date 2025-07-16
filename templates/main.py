# app.py
from flask import Flask, render_template, request, redirect, url_for
import zipfile, xml.etree.ElementTree as ET, io
from datetime import datetime
from werkzeug.utils import secure_filename
from PyPDF2 import PdfReader
from docx import Document
from openpyxl import load_workbook

app = Flask(__name__)

def fmt_date(d):
    if isinstance(d, str):
        try:
            d = datetime.fromisoformat(d)
        except:
            return d
    return d.strftime("%-d %B %Y %H:%M")

def parse_pdf(file):
    reader = PdfReader(file)
    info = reader.metadata or {}
    creator = info.get("/Creator", "—")
    owner = info.get("/Author", "—")
    cd = info.get("/CreationDate","")
    if cd.startswith("D:"):
        try:
            dt = datetime.strptime(cd[2:14], "%Y%m%d%H%M%S")
        except:
            dt = datetime.now()
    else:
        dt = datetime.now()
    text = "\n".join([p.extract_text() or "" for p in reader.pages])
    return creator, owner, fmt_date(dt), text

def parse_docx(file):
    data = file.read()
    creator = owner = "—"
    dt = datetime.now()
    z = zipfile.ZipFile(io.BytesIO(data))
    if "docProps/core.xml" in z.namelist():
        xml_str = z.read("docProps/core.xml")
        ns = {
            "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
            "dc": "http://purl.org/dc/elements/1.1/",
            "dcterms": "http://purl.org/dc/terms/"
        }
        root = ET.fromstring(xml_str)
        c = root.find("dc:creator", ns)
        lm= root.find("cp:lastModifiedBy", ns)
        cr= root.find("dcterms:created", ns)
        if c is not None: creator = c.text
        if lm is not None: owner   = lm.text
        if cr is not None:
            try: dt = datetime.fromisoformat(cr.text)
            except: pass
    doc = Document(io.BytesIO(data))
    paras = [p.text for p in doc.paragraphs if p.text.strip()]
    return creator, owner, fmt_date(dt), "\n".join(paras)

def parse_txt(file):
    text = file.read().decode('utf-8',errors='ignore')
    lines= text.strip().splitlines()
    creator = owner = "—"
    if len(lines)>1 and lines[-1].replace(" ","").isalpha():
        creator = owner = lines.pop()
    return creator, owner, fmt_date(datetime.now()), "\n".join(lines)

def parse_xml(file):
    text = file.read().decode('utf-8',errors='ignore')
    root = ET.fromstring(text)
    creator = root.findtext("creator") or "—"
    owner   = root.findtext("lastModifiedBy") or "—"
    dt_text = root.findtext("created")
    dt = datetime.now()
    if dt_text:
        try: dt = datetime.fromisoformat(dt_text)
        except: pass
    return creator, owner, fmt_date(dt), text

def parse_xlsx(file):
    data = file.read()
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    props = wb.properties
    creator = props.creator or "—"
    owner   = props.lastModifiedBy or "—"
    dt = props.created or datetime.now()
    rows = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        for row in ws.iter_rows(values_only=True):
            rows.append("\t".join(str(c) if c else "" for c in row))
    return creator, owner, fmt_date(dt), "\n".join(rows)

@app.route('/', methods=['GET','POST'])
def index():
    entries = []
    if request.method=='POST':
        files = request.files.getlist('files')
        for f in files:
            ext = f.filename.lower().split('.')[-1]
            if ext=='pdf':    parser=parse_pdf
            elif ext=='docx': parser=parse_docx
            elif ext=='txt':  parser=parse_txt
            elif ext=='xml':  parser=parse_xml
            elif ext=='xlsx': parser=parse_xlsx
            else: continue
            creator, owner, dt, content = parser(f)
            entries.append({
                'name': secure_filename(f.filename),
                'creator': creator,
                'owner': owner,
                'creationDate': dt,
                'content': content
            })
    return render_template('index.html', entries=entries)

if __name__=='__main__':
    app.run(debug=True)
